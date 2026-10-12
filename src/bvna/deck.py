"""Build bvna/deck_bvna.pptx in the same format as the PNS deck.

Every number is read from generated artifacts - bvna/audit/results.json, the
BVNA caption manifest, bvna/audit/novelty_raw.json, and (for the correction
slide) the PNS project's audit/results.json. Anything that cannot be found
renders as red [MISSING] and is listed in a warning. The plain-language
"What this means" boxes are written here and checked by the same contract as
the PNS deck: at most 3 sentences, a word budget, and no banned jargon.
"""
from __future__ import annotations

import json
import logging
from datetime import date
from pathlib import Path
from typing import Any, Final

from pptx import Presentation
from pptx.util import Pt

from src.build_deck import Extraction
from src.build_pptx import (
    ACCENT, AUTHOR_NAME, BODY_PT, INK, MARGIN, MISSING, MUTED, REF_PT, SLIDE_H, SLIDE_W,
    TITLE_PT, FigureSlide, add_headline, add_slide, add_slide_number, add_textbox,
    check_explanation, sentence_count, slide_figure, verify, word_count, write_rich,
)
from src.bvna.config import BVNA_AUDIT, BVNA_FIGURES, BVNA_ROOT, LATEST, YEARS
from src.constants import AUDIT_DIR

LOGGER: Final = logging.getLogger(__name__)


def _today() -> str:
    """'October 12, 2026' on every platform (%-d is not portable to Windows)."""
    today = date.today()
    return f"{today:%B} {today.day}, {today.year}"
DECK_PATH: Final[Path] = BVNA_ROOT / "deck_bvna.pptx"
NOTE_PATH: Final[Path] = BVNA_ROOT / "FIRST_DELIVERABLE_NOTE.md"
FIGURE_STEMS: Final[tuple[str, ...]] = ("figure1", "figure2", "figure3", "figure4", "figure5")

# Closest prior papers, chosen by reading the novelty search results. Only the
# identifiers are chosen here; titles, journals and dates come from the cached
# PubMed records in bvna/audit/novelty_raw.json.
CLOSEST_PMIDS: Final[tuple[tuple[str, str], ...]] = (
    ("37643639", "Health-care use AFTER BVNA in treated patients - outcomes, not access."),
    ("42473615", "BVNA safety from FDA device reports - the BVNA twin of the Cooper PNS paper."),
    ("42804299", "Medicare surgery-centre payment policy - relevant because BVNA is almost all facility-based."),
)


class Facts:
    """Dotted-path reader over the result artifacts that records every miss."""

    def __init__(self, ex: Extraction, **sources: Any) -> None:
        self.ex = ex
        self.sources = sources

    def get(self, path: str, fmt: str = "{}") -> str:
        source, *keys = path.split(".")
        node: Any = self.sources.get(source)
        for key in keys:
            if isinstance(node, dict) and key in node:
                node = node[key]
            elif isinstance(node, dict) and key.isdigit() and int(key) in node:
                node = node[int(key)]
            elif isinstance(node, list) and key.isdigit() and int(key) < len(node):
                node = node[int(key)]
            else:
                node = None
                break
        if node is None:
            self.ex.missing.append(path)
            return MISSING
        try:
            return fmt.format(node)
        except (ValueError, TypeError):
            self.ex.missing.append(f"{path} (format {fmt})")
            return MISSING


def load_facts(ex: Extraction) -> Facts:
    results = json.loads((BVNA_AUDIT / "results.json").read_text())
    captions = json.loads((BVNA_AUDIT / "figure_captions.json").read_text())
    novelty = json.loads((BVNA_AUDIT / "novelty_raw.json").read_text())
    pns = json.loads((AUDIT_DIR / "results.json").read_text())
    # Index the last desert year and the closest papers so templates stay flat.
    results["desert_first"] = results["deserts"][0]
    results["desert_last"] = results["deserts"][-1]
    records = {r["pmid"]: r for q in novelty.values() for r in q["records"]}
    results["closest"] = [records.get(pmid) for pmid, _ in CLOSEST_PMIDS]
    spec = results["specialty"]["64628"]["services_pct"]
    results["specialty_core"] = sum(spec.get(k, 0) for k in ("Pain medicine", "Anesthesiology", "PM&R"))
    results["specialty_surgical"] = sum(spec.get(k, 0) for k in
                                        ("Interventional radiology", "Orthopedics", "Neurosurgery"))
    results["hidden_pct"] = 100 - results["concentration"]["64628"]["share_visible_pct"]
    latest_growth = results["growth"]["by_year"][str(LATEST)]
    results["patients_per_entry"] = (latest_growth["benes_64628"]
                                     / results["concentration"]["64628"]["national_provider_entries_high"])
    g = results["growth"]["by_year"]
    results["rate_fold"] = g[str(YEARS[-1])]["per_100k_64628"] / g[str(YEARS[0])]["per_100k_64628"]
    return Facts(ex, r=results, cap=captions, nov=novelty, pns=pns)


# ------------------------------------------------------------------ slides
def bullets_slide(prs, number: int, title: str, items: list[tuple[str, str]],
                  size: int = 22, gap: int = 18, closing: str | None = None) -> None:
    slide = add_slide(prs)
    add_headline(slide, title)
    _, frame = add_textbox(slide, MARGIN.inches, 1.55, SLIDE_W.inches - 2 * MARGIN.inches,
                           4.95 if closing else 5.55)
    for index, (lead, rest) in enumerate(items):
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        paragraph.space_after = Pt(gap)
        if lead:
            write_rich(paragraph, f"•   {lead}", Pt(size), bold=True, color=INK)
            write_rich(paragraph, f" {rest}" if rest else "", Pt(size), color=INK)
        else:
            write_rich(paragraph, f"•   {rest}", Pt(size), color=INK)
    if closing:
        _, frame = add_textbox(slide, MARGIN.inches, 6.45, SLIDE_W.inches - 2 * MARGIN.inches, 0.6)
        write_rich(frame.paragraphs[0], closing, Pt(21), bold=True, color=ACCENT)
    add_slide_number(slide, number)


def title_slide(prs) -> None:
    slide = add_slide(prs)
    for top, height, text, size, bold, color in (
        (2.25, 1.5, "Who reaches basivertebral nerve ablation?", TITLE_PT, True, INK),
        (3.8, 0.8, "Growth, concentration and access in traditional Medicare, 2022-2024", Pt(24), False, ACCENT),
        (5.0, 0.8, f"{AUTHOR_NAME}   |   {_today()}", BODY_PT, False, MUTED),
    ):
        _, frame = add_textbox(slide, MARGIN.inches, top, SLIDE_W.inches - 2 * MARGIN.inches, height)
        write_rich(frame.paragraphs[0], text, size, bold=bold, color=color)


def figure_slides(f: Facts) -> list[FigureSlide]:
    cap = lambda stem: f.get(f"cap.{stem}.title")
    return [
        FigureSlide(
            stem="figure1", headline=cap("figure1"),
            explanation=(
                f"This chart shows how often Medicare paid for this back procedure each year. "
                f"Use rose from {f.get('r.growth.by_year.2022.per_100k_64628', '{:.0f}')} to "
                f"{f.get('r.growth.by_year.2024.per_100k_64628', '{:.0f}')} procedures for every 100,000 people "
                f"on traditional Medicare between 2022 and 2024. Something this new is spreading fast, so who can reach it matters now."
            ),
            say=(f"{f.get('r.rate_fold', '{:.1f}')}-fold per beneficiary in two years. Counting only doctors' claims "
                 f"it is still {f.get('r.growth_robustness.clinician_fold_change', '{:.1f}')}-fold, so this is not "
                 f"double counting. Extra levels per procedure fell from "
                 f"{f.get('r.growth.by_year.2022.addon_ratio', '{:.2f}')} to "
                 f"{f.get('r.growth.by_year.2024.addon_ratio', '{:.2f}')}."),
            question="Did BVNA exist before 2022?",
            answer=(f"Yes, under facility code {f.get('r.legacy.code')} - it shows only "
                    f"{f.get('r.legacy.services_by_year.2021', '{:,.0f}')} services in 2021 here because hospital "
                    f"facility claims are not in this file, so I start the series at 2022."),
        ),
        FigureSlide(
            stem="figure2", headline=cap("figure2"),
            explanation=(
                f"The bars show how much of each procedure public records can see. For this back procedure it "
                f"is only {f.get('r.concentration.64628.share_visible_pct', '{:.0f}')}%, because most doctors who do "
                f"it treat fewer than 11 Medicare patients a year and are hidden. It is spread across many doctors "
                f"who each do a few, not a small expert group."
            ),
            say=(f"{f.get('r.concentration.64628.national_provider_entries_low', '{:,}')}+ provider entries billed it; "
                 f"only {f.get('r.concentration.64628.visible_clinicians')} clinicians cleared the floor. Among them the "
                 f"top 10% do {f.get('r.concentration.64628.top10_share_pct', '{:.0f}')}%."),
            question="Isn't the low concentration just an artifact of the hidden providers?",
            answer=("Partly - hiding the small providers flattens the visible curve - but the geography file "
                    f"independently shows about {f.get('r.patients_per_entry', '{:.0f}')} Medicare patients per "
                    "provider entry, which is a spread-out pattern on its own."),
        ),
        FigureSlide(
            stem="figure3", headline=cap("figure3"),
            explanation=(
                f"Each bar counts states where almost no one got this procedure. By one count it fell from "
                f"{f.get('r.desert_first.states_zero')} states in 2022 to {f.get('r.desert_last.states_zero')} in 2024, "
                f"though a broader count still finds {f.get('r.desert_last.provider_view_states_sparse')}. The gaps are "
                f"closing quickly, but a few states still have almost no local option."
            ),
            say=(f"Between {f.get('r.desert_last.states_zero')} and "
                 f"{f.get('r.desert_last.provider_view_states_sparse')} states in {LATEST}, covering "
                 f"{f.get('r.desert_last.pct_benes_zero', '{:.1f}')}% to "
                 f"{f.get('r.desert_last.provider_view_pct_benes_sparse', '{:.1f}')}% of beneficiaries."),
            question="Why give a range instead of one number?",
            answer=("Because both files hide something different - one drops small states, the other drops small "
                    "providers - and the honest answer sits between them."),
        ),
        FigureSlide(
            stem="figure4", headline=cap("figure4"),
            explanation=(
                f"Each bar shows which kind of doctor does the procedure. Pain, anesthesia and rehabilitation "
                f"doctors do {f.get('r.specialty_core', '{:.0f}')}% of it, almost the same mix as a common "
                f"nerve-burning treatment for back pain. Spine surgeons barely do it, so access depends on whether "
                f"a nearby pain clinic offers it."
            ),
            say=(f"{f.get('r.specialty_core', '{:.0f}')}% pain, anesthesia or PM&R; spine surgery and IR together "
                 f"{f.get('r.specialty_surgical', '{:.1f}')}%."),
            question="Does this include the doctors you can't see?",
            answer="No - only clinicians above the 11-patient floor, so low-volume adopters could look different.",
        ),
        FigureSlide(
            stem="figure5", headline=cap("figure5"),
            explanation=(
                f"Each panel compares this back procedure with three others on the same scale. It grew "
                f"{f.get('r.growth.comparators.64628.pct_change', '{:.0f}')}% in two years, and "
                f"{f.get('r.setting.64628', '{:.0f}')}% happens in a surgery center or hospital, more than any of the "
                f"others. That makes it harder to reach for people who live far from one."
            ),
            say=(f"Facility share {f.get('r.setting.64628', '{:.0f}')}% vs "
                 f"{f.get('r.setting.62323', '{:.0f}')}% for epidurals; non-metro share "
                 f"{f.get('r.rurality.64628.pct_services_nonmetro', '{:.1f}')}%, lowest of the four."),
            question="Is 'least rural' a population claim?",
            answer="No - it is where the visible doctors practise, not where patients live.",
        ),
    ]


def build(ex: Extraction) -> tuple[Path, list[FigureSlide], list[str], list[int], int]:
    f = load_facts(ex)
    prs = Presentation()
    prs.slide_width, prs.slide_height = SLIDE_W, SLIDE_H

    title_slide(prs)
    bullets_slide(prs, 2, "The question", [
        ("", "BVNA is a newer procedure for vertebrogenic low back pain, with its own billing codes since 2022"),
        ("", "Where is it delivered, by whom, and who lives where there is no access?"),
        ("", "Same pipeline as the PNS deck, using only free public Medicare data"),
        ("", "One method change since that deck - next slide"),
    ], size=24, gap=22)

    pns_visible = f.get("r.concentration.64555.share_visible_pct", "{:.0f}")
    bullets_slide(prs, 3, "A correction to my PNS deck, found while setting this up", [
        ("The file I used sees only part of PNS.",
         f"The provider-level file captures {pns_visible}% of national PNS volume; the geography file shows "
         f"{f.get('r.concentration.64555.national_provider_entries_low', '{:,}')}+ PNS provider entries, not "
         f"{f.get('pns.concentration_billers_only.64555.n_providers')}."),
        ("Some 'providers' were surgery centers.",
         f"{f.get('r.concentration.64555.visible_facility_npis')} of those "
         f"{f.get('pns.concentration_billers_only.64555.n_providers')} were surgery centers billing a facility "
         f"claim for a case a doctor also billed."),
        ("So two claims were overstated:",
         f"'only {f.get('pns.concentration_billers_only.64555.n_providers')} providers' and 'the top 10% do "
         f"{f.get('pns.concentration_common_base.64555.top_decile_share', '{:.0%}')}'."),
        ("Fixed for BVNA:",
         "totals from the geography file, concentration among clinicians only, and the share of volume "
         "visible reported next to every concentration number."),
    ], size=20, gap=14)

    figures = figure_slides(f)
    contract = [p for item in figures for p in check_explanation(item.stem, item.explanation)]
    figure_numbers: list[int] = []
    for offset, item in enumerate(figures):
        number = 4 + offset
        figure_numbers.append(number)
        slide_figure(prs, item, number, ex, figures_dir=BVNA_FIGURES)
        prs.slides[number - 1].notes_slide.notes_text_frame.text = "\n".join([
            f"SAY: {item.say}", "", f"LIKELY QUESTION: {item.question}", f"ANSWER: {item.answer}",
        ])

    n = 4 + len(figures)
    bullets_slide(prs, n, "What surprised me", [
        ("Fewer extra levels per procedure.",
         f"Add-on codes per procedure fell from {f.get('r.growth.by_year.2022.addon_ratio', '{:.2f}')} to "
         f"{f.get('r.growth.by_year.2024.addon_ratio', '{:.2f}')} as volume tripled."),
        ("Most BVNA providers are invisible.",
         f"{f.get('r.concentration.64628.national_provider_entries_low', '{:,}')}+ provider entries, but only "
         f"{f.get('r.concentration.64628.visible_clinicians')} clinicians above the reporting floor."),
        ("It is almost never done in an office:",
         f"{f.get('r.setting.64628', '{:.1f}')}% of services are in a facility."),
        ("Spine surgeons barely do it:",
         f"surgery and interventional radiology together are {f.get('r.specialty_surgical', '{:.1f}')}%."),
        ("The old code barely shows up.",
         f"{f.get('r.legacy.code')} appears only {f.get('r.legacy.services_by_year.2020', '{:,.0f}')} and "
         f"{f.get('r.legacy.services_by_year.2021', '{:,.0f}')} times (2020, 2021): hospital facility claims "
         f"are not in this file."),
    ], size=20, gap=12)

    closest = []
    for index, (pmid, note) in enumerate(CLOSEST_PMIDS):
        closest.append((f"{f.get(f'r.closest.{index}.journal')} {f.get(f'r.closest.{index}.date')[:4]}, "
                        f"PMID {pmid}:", note))
    bullets_slide(prs, n + 1, "What I checked before calling this new", [
        ("PubMed, BVNA total:", f"{f.get('nov.Q1_bvna_any.count')} papers."),
        ("BVNA plus access, Medicare, utilization or disparities:",
         f"{f.get('nov.Q2_bvna_access.count')} papers - judging by title, none studies geographic access or who delivers it."),
    ] + closest, size=19, gap=10,
        closing="Verdict: no BVNA access or provider-distribution study found in PubMed.")

    bullets_slide(prs, n + 2, "Limitations, and which way each one pushes", [
        ("The 11-patient floor", f"hides {f.get('r.hidden_pct', '{:.0f}')}% of BVNA volume. "
         "(Flattens the visible curve, which leans toward my 'spread out' finding; the geography counts back it up.)"),
        ("Traditional Medicare only.", f"Medicare Advantage was {f.get('r.medicare_advantage.national_pct', '{:.0f}')}% "
         f"of Medicare in {LATEST} and over half in {f.get('r.medicare_advantage.states_over_half')} states; its "
         "claims are invisible here. (If it restricts BVNA, real access is worse than this shows.)"),
        ("Surgery-center claims double count totals.",
         f"(Inflates volume; counting doctors only, growth is still "
         f"{f.get('r.growth_robustness.clinician_fold_change', '{:.1f}')}-fold.)"),
        ("Practice address is not the patient's home.", "(Understates how far people travel.)"),
        ("No diagnosis or outcome data.", "I can't see who had vertebrogenic pain and went untreated. (Direction unknown.)"),
    ], size=19, gap=10)

    bullets_slide(prs, n + 3, "What I'd need next", [
        ("Data I don't have:", "Medicare Advantage claims and diagnosis codes, to see who qualified and never got it."),
        ("Natural next study:", "map Medicare's regional coverage policies (public) against where BVNA grew, then "
         "repeat at county level with travel distance."),
        ("What I can contribute:", "screening and extraction, data cleaning, and the statistics - the same "
         "pipeline produced both the PNS and the BVNA decks."),
    ], size=22, gap=22)

    slide = add_slide(prs)
    add_headline(slide, "Data and references")
    _, frame = add_textbox(slide, MARGIN.inches, 1.45, SLIDE_W.inches - 2 * MARGIN.inches, 5.5)
    lines = [
        f"CMS Medicare Physician & Other Practitioners - by Provider and Service; by Geography and Service. "
        f"Data years {YEARS[0]}-{LATEST}. data.cms.gov, free, no login.",
        "CMS Medicare Geographic Variation - by National, State & County (fee-for-service denominators, "
        "Medicare Advantage share).",
        f"Codes verified in the files: 64628, 64629 ({f.get('r.descriptors.0.status')}, "
        f"{f.get('r.descriptors.1.status')}); pre-2022 facility code {f.get('r.legacy.code')}.",
    ] + [
        f"{f.get(f'r.closest.{i}.title')} {f.get(f'r.closest.{i}.journal')}. "
        f"{f.get(f'r.closest.{i}.date')}. PMID {pmid}."
        for i, (pmid, _) in enumerate(CLOSEST_PMIDS)
    ]
    for index, line in enumerate(lines):
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        paragraph.space_after = Pt(9)
        write_rich(paragraph, line, Pt(14), color=INK)
    add_slide_number(slide, n + 4)

    total = n + 4
    prs.save(str(DECK_PATH))
    return DECK_PATH, figures, contract, figure_numbers, total


def write_note(ex: Extraction) -> Path:
    """Short note for the two-week deliverable: growth + concentration + surprises."""
    f = load_facts(ex)
    text = f"""# BVNA - first deliverable note

Growth and concentration charts: `bvna/figures/figure1.png` and `figure2.png`.
Every number below is generated from `bvna/audit/results.json`.

**Codes.** 64628 and 64629 confirmed by their own descriptors in the CMS files for
{YEARS[0]}-{LATEST}. Before 2022 BVNA was billed under facility code
{f.get('r.legacy.code')}, which appears in the physician file only
{f.get('r.legacy.services_by_year.2020', '{:,.0f}')} times in 2020 and
{f.get('r.legacy.services_by_year.2021', '{:,.0f}')} in 2021, because hospital
facility claims are not in this file.

**Growth.** {f.get('r.growth.by_year.2022.per_100k_64628', '{:.1f}')} to
{f.get('r.growth.by_year.2024.per_100k_64628', '{:.1f}')} procedures per 100,000
fee-for-service beneficiaries, {YEARS[0]}-{LATEST} ({f.get('r.rate_fold', '{:.1f}')}-fold).
Counting clinician claims only, volume still rose
{f.get('r.growth_robustness.clinician_fold_change', '{:.1f}')}-fold.

**Concentration.** {f.get('r.concentration.64628.national_provider_entries_low', '{:,}')}+
provider entries billed BVNA in {LATEST}, but only
{f.get('r.concentration.64628.visible_clinicians')} clinicians cleared the
11-patient reporting floor, and the provider file sees
{f.get('r.concentration.64628.share_visible_pct', '{:.0f}')}% of national volume.
Among visible clinicians the top 10% deliver
{f.get('r.concentration.64628.top10_share_pct', '{:.0f}')}%
(Gini {f.get('r.concentration.64628.gini', '{:.2f}')}).

## What surprised me

1. **Levels per procedure fell** from {f.get('r.growth.by_year.2022.addon_ratio', '{:.2f}')}
   to {f.get('r.growth.by_year.2024.addon_ratio', '{:.2f}')} add-on codes per primary code.
2. **BVNA is not concentrated in an expert core.** It is spread across many
   low-volume providers - the opposite of how I framed PNS.
3. **{f.get('r.setting.64628', '{:.1f}')}% is facility-based**, more than PNS, facet RFA or epidurals.
4. **Spine surgeons and IR together do {f.get('r.specialty_surgical', '{:.1f}')}%** of visible BVNA.

## A correction to the PNS deck

The same check showed the provider-level file sees only
{f.get('r.concentration.64555.share_visible_pct', '{:.0f}')}% of national PNS
volume, and {f.get('r.concentration.64555.visible_facility_npis')} of the
{f.get('pns.concentration_billers_only.64555.n_providers')} PNS "providers" were
surgery centers. The PNS provider count and top-10% share were overstated; the
BVNA method above corrects for both.
"""
    NOTE_PATH.write_text(text)
    return NOTE_PATH


def main() -> int:
    ex = Extraction()
    path, figures, contract, figure_numbers, total = build(ex)
    note = write_note(ex)
    LOGGER.info("wrote %s (%d slides, %.2f MB) and %s", path.name, total,
                path.stat().st_size / 1_048_576, note.name)
    problems, explanations = verify(path, total, figure_numbers)
    for label, issues in (("plain-language contract", contract), ("structure", problems)):
        if issues:
            LOGGER.error("%s problems:", label)
            for issue in issues:
                LOGGER.error("  - %s", issue)
        else:
            LOGGER.info("%s OK", label)
    print("\n" + "=" * 74 + f'\n"WHAT THIS MEANS" BOXES ({len(explanations)} figure slides)\n' + "=" * 74)
    for item, text in zip(figures, explanations):
        print(f"\n[{item.stem}] {item.headline}\n  {text}\n  ({word_count(text)} words, "
              f"{sentence_count(text)} sentences)")
    print("=" * 74)
    missing = sorted(set(ex.missing))
    if missing:
        LOGGER.warning("%d MISSING field(s): %s", len(missing), missing)
    else:
        LOGGER.info("no MISSING fields")
    return 1 if (problems or contract or missing) else 0


if __name__ == "__main__":
    logging.basicConfig(level=logging.INFO, format="%(asctime)s  %(levelname)-7s %(message)s",
                        datefmt="%H:%M:%S")
    raise SystemExit(main())
