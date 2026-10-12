"""Build deck.pptx - a widescreen PowerPoint deck for the research meeting.

Every statistic is read programmatically from the generated artifacts via the
extractors in src.build_deck. Nothing numeric is typed by hand. A statistic that
cannot be found renders as red [MISSING] on the slide and is listed in a warning.

The "What this means" text on each figure slide is written fresh here in plain
language - it is deliberately NOT copied from RESULTS.md - with only the numbers
interpolated from the artifacts. A runtime check enforces the sentence count,
the word budget, and the banned-jargon list.
"""
from __future__ import annotations

import logging
import math
import os
import re
from dataclasses import dataclass
from datetime import date
from pathlib import Path
from typing import Callable, Final, Sequence

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Inches, Pt

from src.build_deck import (
    Artifacts,
    Extraction,
    headline_stats,
    limitations,
    references,
)
from src.constants import FIGURES_DIR, PROJECT_ROOT

LOGGER: Final = logging.getLogger(__name__)


def _today() -> str:
    """'October 12, 2026' on every platform (%-d is not portable to Windows)."""
    today = date.today()
    return f"{today:%B} {today.day}, {today.year}"

PPTX_PATH: Final[Path] = PROJECT_ROOT / "deck.pptx"
AUTHOR_NAME: Final[str] = os.environ.get("DECK_AUTHOR", "Armaan Punjani")
MISSING: Final[str] = "[MISSING]"

# ------------------------------------------------------------------ design
SLIDE_W: Final = Inches(13.333)
SLIDE_H: Final = Inches(7.5)
MARGIN: Final = Inches(0.55)

ACCENT: Final = RGBColor(0x0F, 0x5C, 0x8C)      # colorblind-safe deep blue
INK: Final = RGBColor(0x16, 0x19, 0x1C)
MUTED: Final = RGBColor(0x55, 0x60, 0x6A)
BOX_FILL: Final = RGBColor(0xF0, 0xF2, 0xF4)    # light gray explanation box
RED: Final = RGBColor(0xD4, 0x00, 0x00)
WHITE: Final = RGBColor(0xFF, 0xFF, 0xFF)

TITLE_PT: Final = Pt(40)
HEADLINE_PT: Final = Pt(32)
BODY_PT: Final = Pt(22)
EXPLAIN_PT: Final = Pt(18)
EXPLAIN_LABEL_PT: Final = Pt(14)
REF_PT: Final = Pt(11)
SLIDE_NUM_PT: Final = Pt(11)

FONT: Final[str] = "Calibri"

# Figure-slide geometry, in inches. The figure occupies the band between the
# headline and the explanation box, and is fitted inside it without stretching.
# Sized so that the longest headline still fits on two lines at 32pt and the
# longest explanation still fits the box at 18pt, both checked at build time.
HEADLINE_TOP: Final[float] = 0.18
HEADLINE_H: Final[float] = 1.25
FIG_TOP: Final[float] = 1.48
FIG_BAND_H: Final[float] = 3.50
BOX_TOP: Final[float] = 5.15
BOX_H: Final[float] = 1.90

# ------------------------------------------------------ explanation contract
MAX_SENTENCES: Final[int] = 3
MAX_WORDS: Final[int] = 55
BANNED_TERMS: Final[tuple[str, ...]] = (
    "gini", "lorenz", "incidence rate ratio", "negative binomial", "offset",
    "fixed effects", "dual-eligible", "dual eligible", "suppression threshold",
    "ecological", "confounding", "quantile", "coefficient",
)

FIGURE_STEMS: Final[tuple[str, ...]] = (
    "figure1", "figure2", "figure3", "figure4", "figure5", "figure6",
)


@dataclass(frozen=True)
class FigureSlide:
    """One figure slide: headline, image, plain-language box, speaker notes."""

    stem: str
    headline: str
    explanation: str
    say: str
    question: str
    answer: str


# ------------------------------------------------------------------ helpers
def sentence_count(text: str) -> int:
    return len([s for s in re.split(r"(?<=[.!?])\s+", text.strip()) if s])


def word_count(text: str) -> int:
    return len(text.split())


def banned_used(text: str) -> list[str]:
    lowered = text.lower()
    return [term for term in BANNED_TERMS if term in lowered]


def check_explanation(stem: str, text: str) -> list[str]:
    """Enforce the plain-language contract. Returns a list of violations."""
    problems: list[str] = []
    if sentence_count(text) > MAX_SENTENCES:
        problems.append(f"{stem}: {sentence_count(text)} sentences (max {MAX_SENTENCES})")
    if word_count(text) > MAX_WORDS:
        problems.append(f"{stem}: {word_count(text)} words (max {MAX_WORDS})")
    jargon = banned_used(text)
    if jargon:
        problems.append(f"{stem}: banned jargon {jargon}")
    if MISSING in text:
        problems.append(f"{stem}: explanation contains {MISSING}")
    return problems


def strip_paren(value: str) -> str:
    """'0.013 (p = 0.929)' -> '0.013'."""
    return re.sub(r"\s*\(.*?\)\s*", "", value).strip()


def pct_from_ratio(value: str) -> str:
    """'0.022' -> '2%'. Returns MISSING if the value is not a number."""
    try:
        return f"{round(float(value) * 100)}%"
    except (TypeError, ValueError):
        return MISSING


# -------------------------------------------------------- text box plumbing
def add_textbox(slide, left: float, top: float, width: float, height: float):
    box = slide.shapes.add_textbox(Inches(left), Inches(top),
                                   Inches(width), Inches(height))
    frame = box.text_frame
    frame.word_wrap = True
    return box, frame


def style_run(run, size, bold: bool = False, color: RGBColor = INK) -> None:
    run.font.size = size
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = FONT


def write_rich(paragraph, text: str, size, bold: bool = False,
               color: RGBColor = INK) -> None:
    """Write text into a paragraph, rendering any [MISSING] token in red bold."""
    for piece in re.split(rf"({re.escape(MISSING)})", text):
        if not piece:
            continue
        run = paragraph.add_run()
        run.text = piece
        if piece == MISSING:
            style_run(run, size, bold=True, color=RED)
        else:
            style_run(run, size, bold=bold, color=color)


def add_slide(prs: Presentation):
    """A blank white slide."""
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    background = slide.background
    background.fill.solid()
    background.fill.fore_color.rgb = WHITE
    return slide


def add_headline(slide, text: str) -> None:
    _, frame = add_textbox(slide, MARGIN.inches, HEADLINE_TOP,
                           SLIDE_W.inches - 2 * MARGIN.inches, HEADLINE_H)
    frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    write_rich(frame.paragraphs[0], text, HEADLINE_PT, bold=True, color=INK)


def add_slide_number(slide, number: int) -> None:
    _, frame = add_textbox(slide, SLIDE_W.inches - Inches(1.3).inches,
                           SLIDE_H.inches - Inches(0.38).inches, 0.85, 0.3)
    paragraph = frame.paragraphs[0]
    paragraph.alignment = PP_ALIGN.RIGHT
    write_rich(paragraph, str(number), SLIDE_NUM_PT, color=MUTED)


def add_notes(slide, text: str) -> None:
    slide.notes_slide.notes_text_frame.text = text


# ------------------------------------------------- the plain-language boxes
def build_figure_slides(art: Artifacts, ex: Extraction) -> list[FigureSlide]:
    """Compose each figure slide. Explanation prose is authored here, in plain
    words; only the numbers inside it come from the artifacts."""
    stats = headline_stats(art, ex)

    national = art.national
    pns = national[national["hcpcs_code"] == "64555"].sort_values("year")
    if pns.empty:
        ex.missing.append("national.64555")
        rate_now = trough_rate = trough_year = MISSING
    else:
        rate_now = f"{pns.iloc[-1]['services_per_100k']:.0f}"
        trough_rate = f"{pns['services_per_100k'].min():.0f}"
        trough_year = str(int(pns.loc[pns["services_per_100k"].idxmin(), "year"]))

    year = ex.get("pptx.latest_year", stats["latest_year"])
    top10 = ex.get("pptx.top10_pns", stats["top10_pns"])
    states_zero = ex.get("pptx.states_zero", stats["states_zero"])
    pct_sparse = ex.get("pptx.pct_benes_sparse", stats["pct_benes_sparse"])
    r2_pct = ex.get("pptx.pseudo_r2_pct", pct_from_ratio(stats["pseudo_r2_b"]))
    income_r = ex.get("pptx.income_r", strip_paren(stats["income_pearson"]))
    scs_total = ex.get("pptx.scs_total", stats["lit_scs_total"])
    pns_total = ex.get("pptx.pns_total", stats["lit_pns_total"])
    n_pns = ex.get("pptx.pns_providers", stats["pns_providers_parquet"])
    n_epi = ex.get("pptx.epi_providers", stats["epi_providers_parquet"])
    ratio = ex.get("pptx.provider_ratio", stats["provider_ratio"])

    # Headlines come from the caption manifest the figure renderer writes, so
    # the slide headline is literally the sentence printed on the image.
    def headline(stem: str) -> str:
        caption = art.captions.get(stem, {})
        return ex.get(f"pptx.{stem}.headline", caption.get("title"))

    return [
        FigureSlide(
            stem="figure1",
            headline=headline("figure1"),
            explanation=(
                f"This line shows how often each procedure was billed each year. "
                f"Nerve stimulator use almost vanished by {trough_year}, then climbed "
                f"back to {rate_now} procedures for every 100,000 older adults by {year}. "
                f"It is growing fast, so who can actually get one matters now."
            ),
            say=f"{rate_now} per 100,000 in {year}, up from {trough_rate} at the "
                f"{trough_year} low point.",
            question="Why does the line start high in 2015 and then collapse?",
            answer="I don't know, and I say so - the mix of specialties billing it "
                   "is stable the whole time, so it isn't one group handing off to "
                   "another, and I quote growth from the low point with that caveat.",
        ),
        FigureSlide(
            stem="figure2",
            headline=headline("figure2"),
            explanation=(
                f"Each line shows how spread out a procedure is across doctors. "
                f"A small group does almost all the nerve stimulator cases, while "
                f"injections are spread across nearly everyone. The busiest 10% of "
                f"providers do {top10} of stimulator work nationwide, so whole regions "
                f"have nobody offering it."
            ),
            say=f"Top 10% of providers do {top10} of all nerve stimulator volume. "
                f"Only {n_pns} providers billed it nationally, against {n_epi} for "
                f"injections - {ratio}-fold.",
            question="Isn't this just because it's a rare procedure?",
            answer="That's exactly why I rank both procedures across the same group of "
                   "doctors instead of only those who bill each one - and I show the "
                   "other version too, because it flips.",
        ),
        FigureSlide(
            stem="figure3",
            headline=headline("figure3"),
            explanation=(
                f"Each pair of bars is one year. In {year}, {states_zero} states still "
                f"had no doctor billing this procedure at all, and {pct_sparse} of people "
                f"on traditional Medicare lived in a state with almost none. "
                f"The gaps are shrinking, but they have not closed."
            ),
            say=f"{states_zero} states with zero providers in {year}; {pct_sparse} of "
                f"beneficiaries in a state with fewer than three.",
            question="Aren't these gaps closing on their own?",
            answer="Yes, and I lead with that - the finding is that the spread is "
                   "incomplete, not that it isn't happening.",
        ),
        FigureSlide(
            stem="figure4",
            headline=headline("figure4"),
            explanation=(
                f"I checked whether states with more pain doctors also do more nerve "
                f"stimulator cases. They do, but barely: how many pain doctors a state "
                f"has explains only about {r2_pct} of the difference between states. "
                f"Living near a pain doctor does not mean this treatment is within reach."
            ),
            say=f"Counting pain doctors explains about {r2_pct} of the state-to-state "
                f"difference. This one did not go the way I predicted.",
            question="So your hypothesis failed here?",
            answer="Yes - general supply keeps predicting after adjustment, so the "
                   "honest claim is that counting pain doctors is a weak guide to "
                   "access, not a wrong one.",
        ),
        FigureSlide(
            stem="figure5",
            headline=headline("figure5"),
            explanation=(
                f"Each dot is one state, with income rising from left to right. "
                f"I expected richer states to do more nerve stimulator work, but the "
                f"dots scatter with no pattern and the link measures {income_r}, which "
                f"is essentially zero. Money alone does not explain who gets this."
            ),
            say=f"The relationship with state income is {income_r} - a flat line. "
                f"This is a null result and I am not dressing it up.",
            question="Doesn't a null result undercut the whole fairness argument?",
            answer="At the state level, yes - and I think the state is the wrong unit, "
                   "which is why the county version is the next study rather than a "
                   "rescue attempt.",
        ),
        FigureSlide(
            stem="figure6",
            headline=headline("figure6"),
            explanation=(
                f"I counted published studies about who can get each treatment. "
                f"Over ten years there are {scs_total} for spinal cord stimulators and "
                f"{pns_total} for nerve stimulators. Almost nobody has asked who can "
                f"reach the newer one, which is the gap this project fills."
            ),
            say=f"{scs_total} papers on one, {pns_total} on the other, using the same "
                f"search with one word changed.",
            question="Keyword counts don't really prove a gap, do they?",
            answer="Correct - they measure attention, not absence, and I raise that "
                   "myself before anyone raises it for me.",
        ),
    ]


# ---------------------------------------------------------------- slides
def slide_title(prs: Presentation) -> None:
    """Slide 1. No slide number, per the brief."""
    slide = add_slide(prs)
    _, frame = add_textbox(slide, MARGIN.inches, 2.35,
                           SLIDE_W.inches - 2 * MARGIN.inches, 1.5)
    write_rich(frame.paragraphs[0],
               "Who never reaches peripheral nerve stimulation?",
               TITLE_PT, bold=True, color=INK)

    _, frame = add_textbox(slide, MARGIN.inches, 3.85,
                           SLIDE_W.inches - 2 * MARGIN.inches, 0.8)
    write_rich(frame.paragraphs[0],
               "A national analysis using free public Medicare data",
               Pt(24), color=ACCENT)

    _, frame = add_textbox(slide, MARGIN.inches, 5.05,
                           SLIDE_W.inches - 2 * MARGIN.inches, 0.8)
    write_rich(frame.paragraphs[0],
               f"{AUTHOR_NAME}   |   {_today()}",
               BODY_PT, color=MUTED)


def slide_question(prs: Presentation, number: int) -> None:
    """Slide 2. Four large bullets, exactly as specified."""
    slide = add_slide(prs)
    add_headline(slide, "The question, in plain words")

    lines = [
        "Chow et al. 2026: patients who waited longer for insurance approval "
        "got less relief from PNS",
        "But every patient in that study received the device",
        "This project asks the question upstream: who never gets one at all?",
        "Finding: PNS delivery is highly concentrated, so many regions have none",
    ]
    _, frame = add_textbox(slide, MARGIN.inches, 1.75,
                           SLIDE_W.inches - 2 * MARGIN.inches, 4.9)
    for index, line in enumerate(lines):
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        paragraph.space_after = Pt(22)
        write_rich(paragraph, f"•   {line}", Pt(24), color=INK)
    add_slide_number(slide, number)


def slide_figure(prs: Presentation, item: FigureSlide, number: int,
                 ex: Extraction, figures_dir: Path = FIGURES_DIR) -> None:
    """Headline, figure fitted without stretching, explanation box."""
    slide = add_slide(prs)
    add_headline(slide, item.headline)

    path = figures_dir / f"{item.stem}.png"
    if path.exists():
        from PIL import Image

        with Image.open(path) as image:
            aspect = image.size[0] / image.size[1]
        max_w = SLIDE_W.inches - 2 * MARGIN.inches
        height = FIG_BAND_H
        width = height * aspect
        if width > max_w:                 # width-constrained instead
            width = max_w
            height = width / aspect
        left = (SLIDE_W.inches - width) / 2
        top = FIG_TOP + (FIG_BAND_H - height) / 2
        slide.shapes.add_picture(str(path), Inches(left), Inches(top),
                                 Inches(width), Inches(height))
    else:
        ex.missing.append(f"{item.stem}.png")
        _, frame = add_textbox(slide, MARGIN.inches, FIG_TOP,
                               SLIDE_W.inches - 2 * MARGIN.inches, FIG_BAND_H)
        write_rich(frame.paragraphs[0], f"{MISSING} figure {item.stem}.png",
                   Pt(24))

    box = slide.shapes.add_textbox(MARGIN, Inches(BOX_TOP),
                                   Emu(SLIDE_W - 2 * MARGIN), Inches(BOX_H))
    box.fill.solid()
    box.fill.fore_color.rgb = BOX_FILL
    box.line.fill.background()
    frame = box.text_frame
    frame.word_wrap = True
    frame.margin_left = Inches(0.28)
    frame.margin_right = Inches(0.28)
    frame.margin_top = Inches(0.14)
    frame.margin_bottom = Inches(0.14)

    label = frame.paragraphs[0]
    label.space_after = Pt(4)
    write_rich(label, "WHAT THIS MEANS", EXPLAIN_LABEL_PT, bold=True, color=ACCENT)
    body = frame.add_paragraph()
    write_rich(body, item.explanation, EXPLAIN_PT, color=INK)

    add_notes(slide, "\n".join([
        f"SAY: {item.say}",
        "",
        f"LIKELY QUESTION: {item.question}",
        f"ANSWER: {item.answer}",
    ]))
    add_slide_number(slide, number)


def slide_limitations(prs: Presentation, art: Artifacts, ex: Extraction,
                      number: int) -> None:
    """Slide 9. Each limitation with the direction it biases the result."""
    slide = add_slide(prs)
    add_headline(slide, "Limitations")

    items = limitations(art, ex)
    against = [
        item for item in items
        if isinstance(item["direction"], str)
        and ("against" in item["direction"].lower()
             or "toward the null" in item["direction"].lower())
    ]
    toward = [
        item for item in items
        if isinstance(item["direction"], str)
        and "toward my finding" in item["direction"].lower()
    ]

    _, frame = add_textbox(slide, MARGIN.inches, 1.55,
                           SLIDE_W.inches - 2 * MARGIN.inches, 4.6)
    for index, item in enumerate(items):
        direction = item["direction"]
        short = direction.split(".")[0].replace("Direction:", "").strip() \
            if isinstance(direction, str) else MISSING
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        paragraph.space_after = Pt(11)
        write_rich(paragraph, f"•   {item['title']} ", Pt(20), color=INK)
        write_rich(paragraph, f"(biases {short})", Pt(20), color=MUTED)

    closing = (
        f"{len(against)} of these {len(items)} push against my finding. "
        f"{'None' if not toward else str(len(toward))} pushes toward it."
    )
    _, frame = add_textbox(slide, MARGIN.inches, 6.35,
                           SLIDE_W.inches - 2 * MARGIN.inches, 0.7)
    write_rich(frame.paragraphs[0], closing, Pt(22), bold=True, color=ACCENT)
    add_slide_number(slide, number)


def slide_next(prs: Presentation, number: int) -> None:
    """Slide 10. What I'd need next. No numbers here by design."""
    slide = add_slide(prs)
    add_headline(slide, "What I'd need next")
    lines = [
        ("What I don't have: ",
         "patient-level records. I can see which doctors billed, not which "
         "patients were referred, offered the therapy, or turned down."),
        ("The natural next study: ",
         "the same analysis at county level, linked to travel distance and "
         "neighbourhood deprivation, where a real gap would actually show up."),
        ("What I can contribute: ",
         "screening and extraction, data cleaning, and the statistics - "
         "the pipeline behind this deck is one command and it is already written."),
    ]
    _, frame = add_textbox(slide, MARGIN.inches, 1.8,
                           SLIDE_W.inches - 2 * MARGIN.inches, 4.6)
    for index, (lead, rest) in enumerate(lines):
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        paragraph.space_after = Pt(26)
        write_rich(paragraph, f"•   {lead}", Pt(22), bold=True, color=INK)
        write_rich(paragraph, rest, Pt(22), color=INK)
    add_slide_number(slide, number)


def slide_references(prs: Presentation, art: Artifacts, ex: Extraction,
                     number: int) -> None:
    """Slide 11. Only entries that passed the citation audit."""
    slide = add_slide(prs)
    add_headline(slide, "References")

    entries, excluded = references(art, ex)
    _, frame = add_textbox(slide, MARGIN.inches, 1.38,
                           SLIDE_W.inches - 2 * MARGIN.inches, 5.65)
    if not entries:
        write_rich(frame.paragraphs[0], MISSING, REF_PT)
    for index, entry in enumerate(entries):
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        paragraph.space_after = Pt(5)
        text = entry["text"]
        if entry["pmid"] and f"PMID {entry['pmid']}" not in text:
            text = f"{text} PMID {entry['pmid']}."
        write_rich(paragraph, text, REF_PT, color=INK)

    note = frame.add_paragraph()
    note.space_before = Pt(10)
    write_rich(
        note,
        f"All {len(entries)} verified against the publisher record; "
        f"{len(excluded)} excluded as unverifiable.",
        REF_PT, bold=True, color=ACCENT,
    )
    add_slide_number(slide, number)


# ------------------------------------------------------------ verification
# Conservative glyph-width factor for Calibri: average advance width is about
# half the point size, so we assume slightly wider to err toward flagging.
CHAR_WIDTH_FACTOR: Final[float] = 0.52
LINE_HEIGHT_FACTOR: Final[float] = 1.22
EMU_PER_INCH: Final[int] = 914_400


def estimate_overflow(shape) -> tuple[bool, float, float]:
    """Estimate whether a text frame's content exceeds its shape height.

    python-pptx cannot lay out text, so this measures conservatively: it wraps
    each paragraph at the shape's usable width using an average glyph advance
    and sums line heights. Returns (overflows, needed_inches, available_inches).
    """
    frame = shape.text_frame
    usable_w_in = (shape.width - frame.margin_left - frame.margin_right) / EMU_PER_INCH
    available_in = (shape.height - frame.margin_top - frame.margin_bottom) / EMU_PER_INCH
    needed_pt = 0.0
    for paragraph in frame.paragraphs:
        runs = paragraph.runs
        if not runs:
            continue
        size_pt = max((r.font.size.pt for r in runs if r.font.size), default=18.0)
        text = "".join(r.text for r in runs)
        chars_per_line = max(int(usable_w_in * 72 / (size_pt * CHAR_WIDTH_FACTOR)), 1)
        lines = max(math.ceil(len(text) / chars_per_line), 1)
        needed_pt += lines * size_pt * LINE_HEIGHT_FACTOR
        if paragraph.space_after:
            needed_pt += paragraph.space_after.pt
        if paragraph.space_before:
            needed_pt += paragraph.space_before.pt
    needed_in = needed_pt / 72.0
    return needed_in > available_in, needed_in, available_in


def verify(path: Path, expected_slides: int,
           figure_slide_indices: Sequence[int]) -> tuple[list[str], list[str]]:
    """Reopen the deck and check structure. Returns (problems, explanation texts)."""
    prs = Presentation(str(path))
    problems: list[str] = []
    explanations: list[str] = []

    slides = list(prs.slides)
    if len(slides) != expected_slides:
        problems.append(f"expected {expected_slides} slides, found {len(slides)}")

    for index, slide in enumerate(slides, start=1):
        texts = [
            s.text_frame.text.strip()
            for s in slide.shapes
            if s.has_text_frame and s.text_frame.text.strip()
        ]
        pictures = [s for s in slide.shapes if s.shape_type == 13]

        if not texts and not pictures:
            problems.append(f"slide {index} is empty")
        if index in figure_slide_indices:
            if not pictures:
                problems.append(f"slide {index} is a figure slide with no image")
            if not slide.has_notes_slide or not slide.notes_slide.notes_text_frame.text.strip():
                problems.append(f"slide {index} has no speaker notes")
            box = next((s for s in slide.shapes
                        if s.has_text_frame and "WHAT THIS MEANS" in s.text_frame.text), None)
            if box is None:
                problems.append(f"slide {index} has no 'What this means' box")
            else:
                body = box.text_frame.text.split("\n", 1)
                explanations.append(body[1].strip() if len(body) > 1 else "")

        for shape in slide.shapes:
            if not shape.has_text_frame or not shape.text_frame.text.strip():
                continue
            overflows, needed, available = estimate_overflow(shape)
            if overflows:
                problems.append(
                    f"slide {index}: text may overflow "
                    f"({needed:.2f}in needed vs {available:.2f}in available): "
                    f"{shape.text_frame.text[:52]!r}"
                )
            for paragraph in shape.text_frame.paragraphs:
                for run in paragraph.runs:
                    if run.font.size and run.font.size.pt < 11:
                        problems.append(
                            f"slide {index}: {run.font.size.pt}pt text is too small"
                        )
    return problems, explanations


# ------------------------------------------------------------------- main
def main() -> int:
    ex = Extraction()
    art = Artifacts.load()

    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H

    figures = build_figure_slides(art, ex)

    contract_problems: list[str] = []
    for item in figures:
        contract_problems.extend(check_explanation(item.stem, item.explanation))

    slide_title(prs)
    slide_question(prs, 2)
    figure_indices: list[int] = []
    for offset, item in enumerate(figures):
        number = 3 + offset
        figure_indices.append(number)
        slide_figure(prs, item, number, ex)
    slide_limitations(prs, art, ex, 3 + len(figures))
    slide_next(prs, 4 + len(figures))
    slide_references(prs, art, ex, 5 + len(figures))

    total = 5 + len(figures)
    prs.save(str(PPTX_PATH))
    LOGGER.info("wrote %s (%d slides, %.2f MB)", PPTX_PATH.name, total,
                PPTX_PATH.stat().st_size / 1_048_576)

    problems, explanations = verify(PPTX_PATH, total, figure_indices)

    if contract_problems:
        LOGGER.error("plain-language contract violations:")
        for problem in contract_problems:
            LOGGER.error("  - %s", problem)
    else:
        LOGGER.info("plain-language contract OK: <=%d sentences, <=%d words, "
                    "no banned jargon in any explanation", MAX_SENTENCES, MAX_WORDS)

    if problems:
        LOGGER.error("deck.pptx structural problems:")
        for problem in problems:
            LOGGER.error("  - %s", problem)
    else:
        LOGGER.info("deck.pptx verified: %d slides, every figure slide has an "
                    "image, speaker notes, and a fitting text box", total)

    print("\n" + "=" * 74)
    print(f'"WHAT THIS MEANS" BOXES  ({len(explanations)} figure slides)')
    print("=" * 74)
    for item, text in zip(figures, explanations):
        print(f"\n[{item.stem}]  {item.headline}")
        print(f"  {text}")
        print(f"  ({word_count(text)} words, {sentence_count(text)} sentences)")
    print("=" * 74 + "\n")

    unique_missing = sorted(set(ex.missing))
    if unique_missing:
        LOGGER.warning("%d MISSING field(s) rendered as %s:",
                       len(unique_missing), MISSING)
        for name in unique_missing:
            LOGGER.warning("  - %s", name)
    else:
        LOGGER.info("no MISSING fields: every number resolved from an artifact")

    return 1 if (problems or contract_problems) else 0


if __name__ == "__main__":
    logging.basicConfig(level=logging.INFO,
                        format="%(asctime)s  %(levelname)-7s %(message)s",
                        datefmt="%H:%M:%S")
    raise SystemExit(main())
